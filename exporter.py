import io
import datetime

try:
    import docx
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    _DOCX_AVAILABLE = True
except Exception:
    docx = None
    _DOCX_AVAILABLE = False

try:
    from reportlab.lib.pagesizes import letter
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib import colors
    _REPORTLAB_AVAILABLE = True
except Exception:
    _REPORTLAB_AVAILABLE = False


def generate_docx_bytes(summary, original_text="", language_name="English", keywords=None, explanation=""):
    """Generates a Microsoft Word (.docx) binary buffer containing the summary output."""
    if not _DOCX_AVAILABLE:
        doc = docx.Document() if docx else None
        buffer = io.BytesIO()
        if doc:
            doc.add_heading("Text Summarization Output", 0)
            doc.save(buffer)
            buffer.seek(0)
            return buffer.getvalue()
        return b""

    doc = docx.Document()
    
    # Title
    title_p = doc.add_paragraph()
    title_run = title_p.add_run("Text Summarization & Translation Output")
    title_run.font.size = Pt(20)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(30, 41, 59)
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Meta Info
    meta_p = doc.add_paragraph()
    meta_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    time_str = datetime.datetime.now().strftime("%B %d, %Y - %H:%M")
    meta_run = meta_p.add_run(f"Generated on: {time_str} | Language: {language_name}")
    meta_run.font.size = Pt(10)
    meta_run.font.italic = True
    meta_run.font.color.rgb = RGBColor(100, 116, 139)
    
    doc.add_paragraph()

    # Section 1: Summary Output
    h1 = doc.add_heading("1. Summary Result", level=1)
    h1.runs[0].font.color.rgb = RGBColor(15, 23, 42)
    
    summary_p = doc.add_paragraph()
    summary_run = summary_p.add_run(summary or "No summary generated.")
    summary_run.font.size = Pt(11)
    summary_run.font.color.rgb = RGBColor(15, 23, 42)

    # Section 2: Keywords
    if keywords:
        h2 = doc.add_heading("2. Extracted Keywords", level=1)
        h2.runs[0].font.color.rgb = RGBColor(15, 23, 42)
        kw_p = doc.add_paragraph()
        kw_run = kw_p.add_run(", ".join(keywords))
        kw_run.font.size = Pt(10.5)
        kw_run.font.bold = True
        kw_run.font.color.rgb = RGBColor(79, 70, 229)

    # Section 3: Explanation
    if explanation:
        h3 = doc.add_heading("3. Detailed Explanation", level=1)
        h3.runs[0].font.color.rgb = RGBColor(15, 23, 42)
        expl_p = doc.add_paragraph()
        expl_run = expl_p.add_run(explanation)
        expl_run.font.size = Pt(10.5)

    # Section 4: Original Text
    if original_text:
        doc.add_page_break()
        h4 = doc.add_heading("Original Input Text", level=1)
        h4.runs[0].font.color.rgb = RGBColor(71, 85, 105)
        orig_p = doc.add_paragraph()
        orig_run = orig_p.add_run(original_text[:4000] + ("..." if len(original_text) > 4000 else ""))
        orig_run.font.size = Pt(9.5)
        orig_run.font.color.rgb = RGBColor(100, 116, 139)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def generate_pdf_bytes(summary, original_text="", language_name="English", keywords=None, explanation=""):
    """Generates a PDF document binary buffer containing the summary output."""
    buffer = io.BytesIO()

    if not _REPORTLAB_AVAILABLE:
        # Basic text fallback in case reportlab is unavailable
        content = f"SUMMARY ({language_name}):\n\n{summary}\n\nKEYWORDS:\n{', '.join(keywords or [])}\n\nEXPLANATION:\n{explanation}"
        return content.encode('utf-8')

    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        rightMargin=40,
        leftMargin=40,
        topMargin=40,
        bottomMargin=40
    )
    
    styles = getSampleStyleSheet()
    
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontSize=20,
        leading=24,
        textColor=colors.HexColor('#1e293b'),
        alignment=1, # Center
        spaceAfter=6
    )

    meta_style = ParagraphStyle(
        'DocMeta',
        parent=styles['Normal'],
        fontSize=10,
        leading=14,
        textColor=colors.HexColor('#64748b'),
        alignment=1, # Center
        spaceAfter=15
    )

    heading_style = ParagraphStyle(
        'DocHeading',
        parent=styles['Heading2'],
        fontSize=14,
        leading=18,
        textColor=colors.HexColor('#4f46e5'),
        spaceBefore=12,
        spaceAfter=6
    )

    body_style = ParagraphStyle(
        'DocBody',
        parent=styles['Normal'],
        fontSize=10.5,
        leading=15,
        textColor=colors.HexColor('#0f172a'),
        spaceAfter=10
    )

    kw_style = ParagraphStyle(
        'DocKeywords',
        parent=styles['Normal'],
        fontSize=10,
        leading=14,
        textColor=colors.HexColor('#4338ca'),
        spaceAfter=10
    )

    elements = []
    
    elements.append(Paragraph("Text Summarization & Translation Output", title_style))
    time_str = datetime.datetime.now().strftime("%B %d, %Y - %H:%M")
    elements.append(Paragraph(f"Language: <b>{language_name}</b> | Generated on: {time_str}", meta_style))
    elements.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor('#e2e8f0'), spaceAfter=15))

    # Summary section
    elements.append(Paragraph("Summary Output", heading_style))
    safe_summary = (summary or "No summary generated.").replace('\n', '<br/>')
    elements.append(Paragraph(safe_summary, body_style))
    elements.append(Spacer(1, 10))

    # Keywords section
    if keywords:
        elements.append(Paragraph("Extracted Keywords", heading_style))
        elements.append(Paragraph("<b>Keywords:</b> " + ", ".join(keywords), kw_style))
        elements.append(Spacer(1, 10))

    # Explanation section
    if explanation:
        elements.append(Paragraph("Summary Explanation", heading_style))
        safe_expl = explanation.replace('\n', '<br/>')
        elements.append(Paragraph(safe_expl, body_style))
        elements.append(Spacer(1, 10))

    # Original text
    if original_text:
        elements.append(Spacer(1, 15))
        elements.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor('#cbd5e1'), spaceAfter=10))
        elements.append(Paragraph("Original Input Source", heading_style))
        safe_orig = (original_text[:3000] + ("..." if len(original_text) > 3000 else "")).replace('\n', '<br/>')
        elements.append(Paragraph(safe_orig, ParagraphStyle('OrigText', parent=body_style, fontSize=9, leading=13, textColor=colors.HexColor('#475569'))))

    doc.build(elements)
    buffer.seek(0)
    return buffer.getvalue()
